import openpyxl
import os
import sys
import datetime
from pymongo import MongoClient, InsertOne
from dotenv import load_dotenv

sys.stdout.reconfigure(line_buffering=True)

# Load backend environment
load_dotenv(os.path.join(os.path.dirname(__file__), '../../.env'))
mongo_uri = os.environ.get('MONGODB_URI', 'mongodb://127.0.0.1:27017/sld_system')

print(f"[Seed] Connecting to MongoDB: {mongo_uri}...")
client = MongoClient(mongo_uri)
db = client.get_default_database()
if db is None or db.name == 'admin':
    db = client['sld_system']

other_cases_coll = db['othercaselaws']

excel_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../../ALL COURTS DATA.xlsx'))
if not os.path.exists(excel_path):
    print(f"[Error] Excel file not found at: {excel_path}")
    sys.exit(1)

print(f"[Seed] Loading Excel workbook from {excel_path} (this may take ~30 seconds)...")
start_time = datetime.datetime.now()
wb = openpyxl.load_workbook(excel_path)
ws = wb.active
print(f"[Seed] Workbook loaded in {datetime.datetime.now() - start_time}. Total sheet rows: {ws.max_row}")

# Clear existing othercaselaws collection to ensure clean import
print("[Seed] Clearing existing othercaselaws collection...")
other_cases_coll.delete_many({})

batch_size = 2500
batch = []
total_inserted = 0

for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
    s_no = row[0].value
    citation = row[1].value or 'Nil'
    
    # Case No and hyperlink
    case_no_cell = row[2]
    case_no = case_no_cell.value or ''
    case_no_link = ''
    if case_no_cell.hyperlink and case_no_cell.hyperlink.target:
        case_no_link = case_no_cell.hyperlink.target
    
    case_year = row[3].value
    try:
        case_year_num = int(case_year) if case_year is not None else None
    except:
        case_year_num = None

    parties = row[4].value or ''
    order_date_val = row[5].value
    order_date_str = ''
    if isinstance(order_date_val, (datetime.datetime, datetime.date)):
        order_date_str = order_date_val.strftime('%d-%b-%y')
    elif order_date_val is not None:
        order_date_str = str(order_date_val).strip()

    head_notes = row[6].value or ''
    bench = row[7].value or ''
    apex_court = row[8].value or ''
    apex_status = row[9].value or ''
    
    download_cell = row[10] if len(row) > 10 else None
    download_link = ''
    if download_cell:
        if download_cell.hyperlink and download_cell.hyperlink.target:
            download_link = download_cell.hyperlink.target
        elif download_cell.value:
            download_link = str(download_cell.value).strip()
    
    # If download_link is empty but case_no_link is a download url, share it
    if not download_link and case_no_link:
        download_link = case_no_link

    sr_num = None
    if s_no is not None:
        try:
            sr_num = int(s_no)
        except:
            sr_num = row_idx - 1
    else:
        sr_num = row_idx - 1

    doc_seq = row_idx - 1
    doc = {
        'otherCaseId': f"OCL-{str(doc_seq).zfill(6)}",
        'other_case_id': f"OCL-{str(doc_seq).zfill(6)}",
        'srNumber': sr_num,
        'citation': str(citation).strip(),
        'caseNo': str(case_no).strip(),
        'caseNoLink': str(case_no_link).strip(),
        'caseYear': case_year_num,
        'parties': str(parties).strip(),
        'caseTitle': str(parties).strip(),
        'orderDate': order_date_str,
        'judgmentDate': order_date_str,
        'headNotes': str(head_notes).strip() if head_notes else '',
        'bench': str(bench).strip() if bench else '',
        'authorJudge': str(bench).strip() if bench else '',
        'apexCourt': str(apex_court).strip() if apex_court else '',
        'apexStatus': str(apex_status).strip() if apex_status else '',
        'downloadLink': str(download_link).strip(),
        'attachment': str(download_link).strip(),
        'attachmentName': f"{case_no}.pdf" if case_no else 'Judgment.pdf',
        'isDeleted': False,
        'createdAt': datetime.datetime.now(),
        'updatedAt': datetime.datetime.now()
    }

    batch.append(InsertOne(doc))

    if len(batch) >= batch_size:
        other_cases_coll.bulk_write(batch, ordered=False)
        total_inserted += len(batch)
        print(f"[Seed] Inserted {total_inserted} records...")
        batch = []

if batch:
    other_cases_coll.bulk_write(batch, ordered=False)
    total_inserted += len(batch)

print(f"[Seed] Successfully imported {total_inserted} other court cases into MongoDB!")

# Create indexes for search performance
print("[Seed] Creating indexes...")
other_cases_coll.create_index([("caseNo", "text"), ("parties", "text"), ("bench", "text"), ("citation", "text")])
other_cases_coll.create_index("srNumber")
other_cases_coll.create_index("caseYear")
other_cases_coll.create_index("isDeleted")

print(f"[Seed] Complete in {datetime.datetime.now() - start_time}. Database count: {other_cases_coll.count_documents({})}")
